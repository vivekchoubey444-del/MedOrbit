import os, io, json, csv, html, uuid
from pathlib import Path
from flask import Flask, request, render_template, send_file, jsonify
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
OUT = PROJECT / "data/document-studio/outputs"
OUT.mkdir(parents=True, exist_ok=True)
app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 15 * 1024 * 1024

def logo_uri():
    p = os.getenv("MEDORBIT_LOGO_PATH", "")
    candidates = [Path(p)] if p else []
    candidates += [PROJECT / x for x in [
        "assets/images/medorbit-logo.png",
        "assets/images/logo.png",
        "assets/logo.png",
        "public/logo.png"
    ]]
    for f in candidates:
        if f.is_file() and f.suffix.lower() in {".png", ".jpg", ".jpeg"}:
            import base64
            mime = "image/png" if f.suffix.lower() == ".png" else "image/jpeg"
            return f"data:{mime};base64," + base64.b64encode(f.read_bytes()).decode()
    return ""

def extract(upload):
    name = Path(upload.filename or "document.txt").name
    ext = Path(name).suffix.lower()
    allowed = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".pdf", ".docx", ".xlsx"}
    if ext not in allowed:
        raise ValueError("Supported formats: TXT, MD, CSV, JSON, HTML, PDF, DOCX, XLSX.")
    raw = upload.read()
    if not raw:
        raise ValueError("The uploaded file is empty.")
    title = Path(name).stem.replace("_", " ").replace("-", " ")
    lines, tables = [], []
    if ext in {".txt", ".md"}:
        lines = raw.decode("utf-8-sig", "replace").splitlines()
    elif ext == ".json":
        lines = json.dumps(json.loads(raw.decode("utf-8-sig")), ensure_ascii=False, indent=2).splitlines()
    elif ext in {".html", ".htm"}:
        from html.parser import HTMLParser
        class Extractor(HTMLParser):
            def __init__(self): super().__init__(); self.parts=[]; self.skip=0
            def handle_starttag(self, tag, attrs):
                if tag in {"script","style"}: self.skip += 1
                if tag in {"p","div","br","li","h1","h2","tr"}: self.parts.append("\n")
            def handle_endtag(self, tag):
                if tag in {"script","style"} and self.skip: self.skip -= 1
            def handle_data(self, data):
                if not self.skip: self.parts.append(data)
        parser = Extractor()
        parser.feed(raw.decode("utf-8", "replace"))
        lines = "".join(parser.parts).splitlines()
    elif ext == ".csv":
        tables = list(csv.reader(io.StringIO(raw.decode("utf-8-sig", "replace"))))
    elif ext == ".pdf":
        for i, page in enumerate(PdfReader(io.BytesIO(raw)).pages, 1):
            lines += [f"Page {i}", *(page.extract_text() or "[No extractable text]").splitlines(), ""]
    elif ext == ".docx":
        doc = Document(io.BytesIO(raw))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    elif ext == ".xlsx":
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        for sheet in book.worksheets:
            lines.append("Worksheet: " + sheet.title)
            rows = [[str(v) if v is not None else "" for v in row] for row in sheet.iter_rows(values_only=True)]
            if rows: tables.append(rows)
        book.close()
    return title, name, lines, tables

def esc(value):
    return html.escape(str(value), quote=True)

def html_doc(title, source, lines, tables):
    body = "".join("<p>"+esc(x)+"</p>" if x.strip() else "<div class='space'></div>" for x in lines)
    for rows in tables:
        if not rows: continue
        body += "<div class='scroll'><table>"
        for i, row in enumerate(rows):
            tag = "th" if i == 0 else "td"
            body += "<tr>" + "".join(f"<{tag}>{esc(v)}</{tag}>" for v in row) + "</tr>"
        body += "</table></div>"
    return render_template("document.html", title=title, source=source, body=body, logo=logo_uri())

def create_pdf(path, title, source, lines, tables):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="MOtitle", parent=styles["Title"],
        fontName="Helvetica-Bold", fontSize=20, leading=25,
        textColor=colors.HexColor("#263e72"), alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="MObody", parent=styles["BodyText"],
        fontSize=9, leading=13, textColor=colors.HexColor("#263238"), spaceAfter=4))
    story = [Paragraph("MEDORBIT", styles["MOtitle"]),
             Paragraph(esc(title), styles["Heading1"]),
             Paragraph("Source: "+esc(source), styles["Italic"]), Spacer(1, 6*mm)]
    for line in lines:
        story.append(Paragraph(esc(line) if line.strip() else " ", styles["MObody"]))
    for rows in tables:
        if not rows: continue
        data = [[Paragraph(esc(v)[:2500], styles["MObody"]) for v in row] for row in rows]
        cols = max(map(len, data))
        data = [r + [""]*(cols-len(r)) for r in data]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#263e72")),
            ("TEXTCOLOR",(0,0),(-1,0),colors.white),
            ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#cbd5e1")),
            ("VALIGN",(0,0),(-1,-1),"TOP")
        ]))
        story += [Spacer(1,4*mm), t, Spacer(1,4*mm)]
    def footer(canvas, doc):
        canvas.saveState()
        w,h=A4
        canvas.setStrokeColor(colors.HexColor("#d6b46a"))
        canvas.line(19*mm,h-15*mm,w-19*mm,h-15*mm)
        canvas.setFillColor(colors.HexColor("#263e72"))
        canvas.setFont("Helvetica-Bold",8)
        canvas.drawString(19*mm,10*mm,"MedOrbit | Education & Career Guidance")
        canvas.setFont("Helvetica",8)
        canvas.drawRightString(w-19*mm,10*mm,f"Page {doc.page}")
        canvas.restoreState()
    SimpleDocTemplate(str(path),pagesize=A4,leftMargin=19*mm,rightMargin=19*mm,
        topMargin=23*mm,bottomMargin=22*mm,title=title,author="MedOrbit").build(
        story,onFirstPage=footer,onLaterPages=footer)

@app.get("/")
def home():
    return render_template("index.html", logo=logo_uri())

@app.get("/health")
def health():
    return jsonify(status="ok", service="MedOrbit Document Studio")

@app.post("/api/convert")
def convert():
    try:
        upload=request.files.get("file")
        if not upload or not upload.filename:
            return jsonify(error="Choose a file first."),400
        title, source, lines, tables=extract(upload)
        fmt=request.form.get("format","pdf").lower()
        if fmt not in {"pdf","html"}:
            return jsonify(error="Format must be pdf or html."),400
        path=OUT/f"medorbit-{uuid.uuid4().hex[:12]}.{fmt}"
        if fmt=="pdf":
            create_pdf(path,title,source,lines,tables)
            mime="application/pdf"
        else:
            path.write_text(html_doc(title,source,lines,tables),encoding="utf-8")
            mime="text/html"
        return send_file(path,as_attachment=True,download_name=f"MedOrbit-{path.stem}.{fmt}",mimetype=mime)
    except ValueError as e:
        return jsonify(error=str(e)),400
    except Exception as e:
        app.logger.exception("Conversion failed")
        return jsonify(error=f"Conversion failed: {type(e).__name__}: {e}"),500

@app.errorhandler(413)
def too_large(_):
    return jsonify(error="Maximum upload size is 15 MB."),413

if __name__=="__main__":
    app.run(host="0.0.0.0",port=int(os.getenv("PORT","4170")),debug=False)
